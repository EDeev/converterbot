import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.shared import OxmlElement, qn
from docx.shared import Cm, Inches, Pt, RGBColor

# Структурные элементы по ГОСТ 7.32-2017 — заголовки без номера
STRUCTURAL_HEADINGS = re.compile(
    r"^(реферат|содержание|оглавление|введение|заключение|список\s+(использованных\s+)?(литературы|источников)"
    r"|библиография|bibliography|references|приложени[ея].*|термины\s+и\s+определения"
    r"|перечень\s+сокращений.*|определения|обозначения\s+и\s+сокращения)$",
    re.IGNORECASE,
)
BIBLIOGRAPHY_HEADING = re.compile(
    r"^(список\s+(использованных\s+)?(литературы|источников)|библиография|bibliography|references)$",
    re.IGNORECASE,
)


def set_style_font(style, font_name):
    """Шрифт стиля для всех письменностей. Встроенные стили Word (заголовки) задают шрифт темы
    (asciiTheme и т. п.), который перекрывает font.name — без очистки заголовки выходят в Calibri"""
    style.font.name = font_name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), font_name)


def add_page_field(paragraph):
    """Поле PAGE — номер страницы, который Word подставляет сам"""
    run = paragraph.add_run()
    begin, instr, end = OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    return run


class DocumentSettings:
    """Настройки форматирования документа с поддержкой ГОСТ"""
    def __init__(self):
        # Базовые настройки текста
        self.font_name = "Times New Roman"
        self.font_size = 14  # основной текст
        self.line_spacing = 1.5
        self.justify_text = True
        self.paragraph_spacing = 6
        self.text_color = (0, 0, 0)
        self.paragraph_indent = 1.25
        
        # Отступы от полей документа в сантиметрах (ГОСТ 7.32-2017)
        self.margin_top = 2.0
        self.margin_bottom = 2.0
        self.margin_left = 3.0    # увеличено для переплета
        self.margin_right = 1.5
        
        # Настройки шрифтов заголовков по ГОСТ
        self.heading1_font_size = 16  # Заголовки глав
        self.heading2_font_size = 14  # Заголовки разделов  
        self.heading3_font_size = 14  # Подзаголовки
        self.heading4_font_size = 14
        self.heading5_font_size = 12
        self.heading6_font_size = 12
        self.footnote_font_size = 10
        
        # Интервалы заголовков по ГОСТ
        self.heading_spacing_before = 12  # пт
        self.heading_spacing_after = 6    # пт
        self.paragraph_spacing_before = 0 # пт
        
        # Нумерация страниц
        self.page_numbering = True
        self.page_number_position = "bottom_center"  # top_right, bottom_center, bottom_right
        self.page_number_start = 1
        self.exclude_title_page_numbering = True
        
        # Автонумерация заголовков
        self.auto_numbering_headings = False
        self.numbering_format = "decimal"  # "decimal" (1.1.1) или "simple" (1)
        
        # Дополнительные ГОСТ настройки
        self.bibliography_style = "gost"
        self.table_caption_position = "above"  # above, below
        self.figure_caption_position = "below"



class MarkdownToDocxConverter:
    """Конвертер Markdown в DOCX с поддержкой ГОСТ"""
    
    def __init__(self, settings: DocumentSettings = None):
        self.settings = settings or DocumentSettings()
        self.doc = Document()
        
        # Счетчики для автонумерации
        self.heading_counters = [0] * 6  # для 6 уровней заголовков
        self.footnote_counter = 0
        self.table_counter = 0
        self.figure_counter = 0
        
        self.setup_document_margins()
        self.setup_page_numbering()
        self.setup_styles()
        
    def setup_document_margins(self):
        """Настройка отступов от полей документа по ГОСТ"""
        sections = self.doc.sections
        for section in sections:
            section.top_margin = Cm(self.settings.margin_top)
            section.bottom_margin = Cm(self.settings.margin_bottom)
            section.left_margin = Cm(self.settings.margin_left)
            section.right_margin = Cm(self.settings.margin_right)
            
    def setup_page_numbering(self):
        """Настройка нумерации страниц согласно ГОСТ"""
        if not self.settings.page_numbering:
            return
            
        section = self.doc.sections[0]

        # Создание колонтитула для нумерации (раньше колонтитул создавался, но поле номера
        # страницы в него не добавлялось — номеров в документе не было)
        if self.settings.page_number_position == "top_right":
            para = section.header.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            para = section.footer.paragraphs[0]
            para.alignment = (WD_ALIGN_PARAGRAPH.RIGHT if self.settings.page_number_position == "bottom_right"
                              else WD_ALIGN_PARAGRAPH.CENTER)
        para.paragraph_format.first_line_indent = Cm(0)
        run = add_page_field(para)
        run.font.name = self.settings.font_name
        run.font.size = Pt(self.settings.font_size)

        # титульный лист без номера: у первой страницы свой, пустой колонтитул
        if self.settings.exclude_title_page_numbering:
            section.different_first_page_header_footer = True
    
    def setup_styles(self):
        """Настройка стилей документа в соответствии с ГОСТ"""
        styles = self.doc.styles
        
        # Настройка базового стиля
        normal_style = styles['Normal']
        normal_font = normal_style.font
        set_style_font(normal_style, self.settings.font_name)
        normal_font.size = Pt(self.settings.font_size)
        normal_font.color.rgb = RGBColor(*self.settings.text_color)
        
        normal_paragraph = normal_style.paragraph_format
        normal_paragraph.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        normal_paragraph.line_spacing = self.settings.line_spacing
        normal_paragraph.space_after = Pt(self.settings.paragraph_spacing)
        normal_paragraph.space_before = Pt(self.settings.paragraph_spacing_before)
        normal_paragraph.first_line_indent = Cm(self.settings.paragraph_indent)
        
        if self.settings.justify_text:
            normal_paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
        # Настройка стилей заголовков с дифференцированными размерами
        heading_sizes = [
            self.settings.heading1_font_size,
            self.settings.heading2_font_size,
            self.settings.heading3_font_size,
            self.settings.heading4_font_size,
            self.settings.heading5_font_size,
            self.settings.heading6_font_size
        ]
        
        for i in range(1, 7):
            heading_style_name = f'Heading {i}'
            if heading_style_name in [s.name for s in styles]:
                heading_style = styles[heading_style_name]
            else:
                heading_style = styles.add_style(heading_style_name, WD_STYLE_TYPE.PARAGRAPH)
                
            heading_font = heading_style.font
            set_style_font(heading_style, self.settings.font_name)
            heading_font.size = Pt(heading_sizes[i-1])  # используем соответствующий размер
            heading_font.bold = True
            heading_font.italic = False
            heading_font.color.rgb = RGBColor(*self.settings.text_color)
            
            heading_paragraph = heading_style.paragraph_format
            heading_paragraph.space_before = Pt(self.settings.heading_spacing_before)
            heading_paragraph.space_after = Pt(self.settings.heading_spacing_after)
            
            # Заголовки 1 и 2 уровня по центру (ГОСТ), остальные с отступом
            if i <= 2:
                heading_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                heading_paragraph.first_line_indent = Cm(0)
            else:
                if self.settings.justify_text:
                    heading_paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                heading_paragraph.first_line_indent = Cm(self.settings.paragraph_indent)
        
        # Стиль для сносок
        try:
            footnote_style = styles.add_style('Footnote', WD_STYLE_TYPE.PARAGRAPH)
            footnote_font = footnote_style.font
            footnote_font.name = self.settings.font_name
            footnote_font.size = Pt(self.settings.footnote_font_size)
            footnote_font.color.rgb = RGBColor(*self.settings.text_color)
            
            footnote_paragraph = footnote_style.paragraph_format
            footnote_paragraph.space_before = Pt(3)
            footnote_paragraph.space_after = Pt(3)
            footnote_paragraph.first_line_indent = Cm(0.5)
        except ValueError:  # стиль уже есть
            pass
            
        # Стиль для кода (без изменений)
        try:
            code_style = styles.add_style('Code', WD_STYLE_TYPE.CHARACTER)
            code_font = code_style.font
            set_style_font(code_style, 'Courier New')
            code_font.size = Pt(self.settings.font_size)
            code_font.color.rgb = RGBColor(*self.settings.text_color)
        except ValueError:  # стиль уже есть
            pass
            
        # Стиль для блоков кода
        try:
            code_block_style = styles.add_style('Code Block', WD_STYLE_TYPE.PARAGRAPH)
            code_block_font = code_block_style.font
            set_style_font(code_block_style, 'Courier New')
            code_block_font.size = Pt(self.settings.font_size)
            code_block_font.color.rgb = RGBColor(*self.settings.text_color)
            
            code_block_paragraph = code_block_style.paragraph_format
            code_block_paragraph.left_indent = Inches(0.5)
            code_block_paragraph.first_line_indent = Cm(0)  # без отступа первой строки для кода
            code_block_paragraph.space_before = Pt(6)
            code_block_paragraph.space_after = Pt(6)
        except ValueError:  # стиль уже есть
            pass
            
        # Стиль для подписей к таблицам и рисункам
        # «Caption» уже есть во встроенном шаблоне (синий, 9 пт) — настраиваем его, а не создаём
        caption_style = (styles['Caption'] if 'Caption' in [s.name for s in styles]
                         else styles.add_style('Caption', WD_STYLE_TYPE.PARAGRAPH))
        caption_font = caption_style.font
        set_style_font(caption_style, self.settings.font_name)
        caption_font.size = Pt(self.settings.font_size - 2)  # меньше основного текста
        caption_font.bold = False
        caption_font.italic = False
        caption_font.color.rgb = RGBColor(*self.settings.text_color)

        caption_paragraph = caption_style.paragraph_format
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph.first_line_indent = Cm(0)
        caption_paragraph.space_before = Pt(6)
        caption_paragraph.space_after = Pt(6)
    
    def generate_heading_number(self, level: int) -> str:
        """Генерация номера заголовка согласно настройкам автонумерации"""
        if not self.settings.auto_numbering_headings:
            return ""
        
        # Обновляем счетчик текущего уровня
        self.heading_counters[level - 1] += 1
        
        # Обнуляем счетчики всех нижестоящих уровней
        for i in range(level, 6):
            self.heading_counters[i] = 0
        
        if self.settings.numbering_format == "simple":
            return f"{self.heading_counters[level - 1]}. "
        else:  # decimal
            # Формируем иерархическую нумерацию
            numbers = []
            for i in range(level):
                if self.heading_counters[i] > 0:
                    numbers.append(str(self.heading_counters[i]))
            return ".".join(numbers) + ". " if numbers else ""
    
    def parse_markdown_file(self, file_path: str):
        """Чтение и парсинг Markdown файла"""
        try:
            with open(file_path, 'r', encoding='utf-8-sig') as file:
                content = file.read()
            return content
        except Exception as e:
            raise Exception(f"Ошибка чтения файла: {e}") from e
    
    def add_text_run_with_color(self, paragraph, text, bold=False, italic=False, code_style=False):
        """Добавление текста с настройкой цвета"""
        run = paragraph.add_run(text)
        run.font.color.rgb = RGBColor(*self.settings.text_color)
        
        if bold:
            run.font.bold = True
        if italic:
            run.font.italic = True
        if code_style:
            run.style = 'Code'
        
        return run
    
    def process_text_formatting(self, text: str, paragraph):
        """Обработка форматирования текста включая сноски [^1]"""
        # Обработка сносок
        footnote_pattern = r'\[\^(\d+)\]'
        footnotes = re.findall(footnote_pattern, text)
        
        # Заменяем сноски на верхние индексы
        for footnote_num in footnotes:
            text = re.sub(rf'\[\^{footnote_num}\]', f'{{FOOTNOTE_{footnote_num}}}', text)
        
        # Разбор текста на части с различным форматированием
        parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\{FOOTNOTE_\d+\})', text)
        
        for part in parts:
            if not part:
                continue
            
            if part.startswith('**') and part.endswith('**'):
                # Жирный текст
                self.add_text_run_with_color(paragraph, part[2:-2], bold=True)
            elif part.startswith('*') and part.endswith('*'):
                # Курсив
                self.add_text_run_with_color(paragraph, part[1:-1], italic=True)
            elif part.startswith('`') and part.endswith('`'):
                # Инлайн код
                self.add_text_run_with_color(paragraph, part[1:-1], code_style=True)
            elif part.startswith('{FOOTNOTE_') and part.endswith('}'):
                # Сноска - добавляем как верхний индекс
                footnote_num = re.search(r'FOOTNOTE_(\d+)', part).group(1)
                run = self.add_text_run_with_color(paragraph, footnote_num)
                run.font.superscript = True
            else:
                # Обычный текст
                self.add_text_run_with_color(paragraph, part)
    
    LIST_ITEM = re.compile(r'^(\s*)([-*+]|\d+[.)])\s+(.*)$')

    def process_list(self, lines: list, start_idx: int):
        """Обработка списков с правильным форматированием по ГОСТ.
        Маркер — тире, нумерация своя для каждого списка и уровня (раньше стиль List Bullet давал
        второй маркер, вложенность терялась, а List Number продолжал счёт из предыдущего списка)"""
        i = start_idx
        counters = {}

        while i < len(lines):
            match = self.LIST_ITEM.match(lines[i].expandtabs(4))
            if not match:
                if lines[i].strip() == '' and i + 1 < len(lines) and self.LIST_ITEM.match(lines[i + 1].expandtabs(4)):
                    i += 1
                    continue
                break

            indent, marker, text = match.groups()
            level = min(len(indent) // 2, 3)
            for deeper in [lvl for lvl in counters if lvl > level]:
                del counters[deeper]

            if marker[0].isdigit():
                counters[level] = counters.get(level, 0) + 1
                prefix = f"{counters[level]}) "
            else:
                counters.pop(level, None)
                prefix = "– "  # тире вместо точек (ГОСТ)

            paragraph = self.doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(level * 0.75)
            paragraph.paragraph_format.first_line_indent = Cm(self.settings.paragraph_indent)
            self.add_text_run_with_color(paragraph, prefix)
            self.process_text_formatting(text, paragraph)
            i += 1

        return i - 1

    TABLE_SEPARATOR = re.compile(r'^\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)*\|?$')

    @staticmethod
    def split_row(line: str) -> list:
        """Ячейки строки таблицы; внешние «|» необязательны"""
        line = line.strip()
        if line.startswith('|'):
            line = line[1:]
        if line.endswith('|'):
            line = line[:-1]
        return [cell.strip() for cell in line.split('|')]

    def process_table(self, lines: list, start_idx: int):
        """Обработка таблиц с подписями согласно ГОСТ"""
        i = start_idx
        table_lines = []
        
        # таблица заканчивается на пустой строке — иначе две таблицы подряд сливались в одну
        while i < len(lines):
            line = lines[i].strip()
            if '|' in line:
                table_lines.append(line)
            else:
                break
            i += 1

        if len(table_lines) < 2 or not self.TABLE_SEPARATOR.match(table_lines[1]):
            # не таблица, а строка с «|» — обычный абзац
            paragraph = self.doc.add_paragraph()
            self.process_text_formatting(lines[start_idx].strip(), paragraph)
            return start_idx
        
        # Добавляем подпись к таблице (если настроено)
        if self.settings.table_caption_position == "above":
            self.table_counter += 1
            caption_para = self.doc.add_paragraph()
            caption_para.style = 'Caption'
            caption_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            caption_para.add_run(f"Таблица {self.table_counter}")
        
        # Парсинг и создание таблицы
        headers = self.split_row(table_lines[0])
        data_lines = table_lines[2:] if len(table_lines) > 2 else []
        
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = 'Table Grid'
        
        # Заполнение заголовков
        header_row = table.rows[0]
        for idx, header in enumerate(headers):
            cell = header_row.cells[idx]
            cell.text = header
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.bold = True
        
        # Заполнение данных
        for line in data_lines:
            row_data = self.split_row(line)
            row = table.add_row()
            for idx, cell_data in enumerate(row_data):
                if idx < len(row.cells):
                    row.cells[idx].text = cell_data
                    # Выравнивание по центру для всех ячеек (ГОСТ)
                    for paragraph in row.cells[idx].paragraphs:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # в ячейках — без абзацного отступа и с одинарным интервалом, иначе текст смещён,
        # а строки получаются высокими
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    fmt = paragraph.paragraph_format
                    fmt.first_line_indent = Cm(0)
                    fmt.space_before = Pt(0)
                    fmt.space_after = Pt(0)
                    fmt.line_spacing = 1.0

        # Подпись снизу (если настроено)
        if self.settings.table_caption_position == "below":
            self.table_counter += 1
            caption_para = self.doc.add_paragraph()
            caption_para.style = 'Caption'
            caption_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            caption_para.add_run(f"Таблица {self.table_counter}")
        
        return i - 1
    
    def process_code_block(self, lines: list, start_idx: int):
        """Обработка блоков кода"""
        i = start_idx + 1
        code_lines = []
        
        while i < len(lines):
            line = lines[i]
            if line.strip().startswith('```'):
                break
            code_lines.append(line)
            i += 1
        
        code_paragraph = self.doc.add_paragraph()
        code_paragraph.style = 'Code Block'
        code_paragraph.add_run('\n'.join(code_lines))
        
        return i
    
    def add_footnote_definition(self, footnote_num: str, footnote_text: str):
        """Добавление определения сноски в конец документа"""
        footnote_para = self.doc.add_paragraph()
        footnote_para.style = 'Footnote'
        
        # Номер сноски как верхний индекс
        footnote_run = footnote_para.add_run(footnote_num)
        footnote_run.font.superscript = True
        
        # Текст сноски
        footnote_para.add_run(f" {footnote_text}")
    
    def process_bibliography(self, lines: list, start_idx: int):
        """Обработка списка литературы в стиле ГОСТ"""
        i = start_idx
        bib_items = []
        
        # Поиск элементов библиографии
        while i < len(lines):
            line = lines[i].strip()
            if re.match(r'^(\d+[.)]|[-*+])\s', line):
                bib_text = re.sub(r'^(\d+[.)]|[-*+])\s', '', line)
                bib_items.append(bib_text)
            elif line == '':
                i += 1
                continue
            else:
                break
            i += 1
        
        if bib_items:
            # Элементы библиографии (заголовок уже добавлен в convert)
            for idx, item in enumerate(bib_items, 1):
                bib_para = self.doc.add_paragraph()
                bib_para.paragraph_format.first_line_indent = Cm(0)
                bib_para.paragraph_format.left_indent = Cm(1)
                bib_para.add_run(f"{idx}. {item}")
        
        return i - 1
    
    def convert(self, md_file_path: str, output_path: str = None):
        """Основной метод конвертации с поддержкой ГОСТ"""
        if not output_path:
            md_path = Path(md_file_path)
            output_path = md_path.with_suffix('.docx')
        
        content = self.parse_markdown_file(md_file_path)
        lines = content.split('\n')
        
        # Сбор сносок для обработки в конце
        footnote_definitions = {}
        
        i = 0
        while i < len(lines):
            line = lines[i]
            stripped_line = line.strip()
            
            if not stripped_line:
                i += 1
                continue
            
            # Обработка определений сносок [^1]: текст сноски
            footnote_def_match = re.match(r'^\[\^(\d+)\]:\s*(.+)', stripped_line)
            if footnote_def_match:
                footnote_num = footnote_def_match.group(1)
                footnote_text = footnote_def_match.group(2)
                footnote_definitions[footnote_num] = footnote_text
                i += 1
                continue
            
            # Заголовки с автонумерацией
            if stripped_line.startswith('#'):
                match = re.match(r'^(#{1,6})\s+(.+)', stripped_line)
                if match:
                    level = len(match.group(1))
                    title = match.group(2).strip()
                    structural = bool(STRUCTURAL_HEADINGS.match(title.rstrip(':')))

                    # Разрыв страницы перед заголовком 2 уровня
                    if level == 2:
                        self.doc.add_page_break()

                    heading = self.doc.add_paragraph()
                    heading.style = f'Heading {level}'

                    if structural:
                        # структурные элементы (введение, заключение, список литературы...) по ГОСТ
                        # не нумеруются, пишутся прописными и по центру
                        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        heading.paragraph_format.first_line_indent = Cm(0)
                        self.process_text_formatting(title.upper(), heading)
                        if BIBLIOGRAPHY_HEADING.match(title.rstrip(':')):
                            i = self.process_bibliography(lines, i + 1)
                    else:
                        # Добавляем автонумерацию
                        heading_number = self.generate_heading_number(level)
                        self.process_text_formatting(heading_number + title, heading)
            
            # Блоки кода
            elif stripped_line.startswith('```'):
                i = self.process_code_block(lines, i)
            
            # Списки
            elif self.LIST_ITEM.match(line.expandtabs(4)):
                i = self.process_list(lines, i)

            # Таблицы
            elif '|' in stripped_line:
                i = self.process_table(lines, i)
            
            # Цитаты
            elif stripped_line.startswith('>'):
                quote_text = re.sub(r'^>\s?', '', stripped_line)
                quote_paragraph = self.doc.add_paragraph()
                quote_paragraph.paragraph_format.left_indent = Inches(0.5)
                quote_paragraph.paragraph_format.right_indent = Inches(0.5)
                self.process_text_formatting(quote_text, quote_paragraph)
                
                for run in quote_paragraph.runs:
                    run.font.italic = True
            
            # Горизонтальные линии
            elif stripped_line in ['---', '***', '___']:
                hr_paragraph = self.doc.add_paragraph()
                hr_paragraph.add_run('_' * 50)
                hr_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Обычные абзацы
            else:
                paragraph = self.doc.add_paragraph()
                self.process_text_formatting(stripped_line, paragraph)
            
            i += 1
        
        # Добавление сносок в конец документа
        if footnote_definitions:
            # Разделительная линия
            self.doc.add_paragraph().add_run('_' * 50)
            
            for footnote_num in sorted(footnote_definitions.keys(), key=int):
                self.add_footnote_definition(footnote_num, footnote_definitions[footnote_num])
        
        self.doc.save(output_path)
        return output_path

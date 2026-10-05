import zipfile

from docx import Document

from md2gost import DocumentSettings, MarkdownToDocxConverter
from md2gost.cli import main as cli_main

SAMPLE = """# Отчёт о практике

## Введение

Абзац с **жирным**, *курсивом*, `кодом` и сноской[^1].

### Цели

- Первый пункт
- Второй пункт
  - Вложенный пункт

1. Раз
2. Два

Второй список:

1. Снова один
2. Снова два

| Параметр | Значение |
|---|---|
| A | 1 |

Строка с | вертикальной чертой, но не таблица.

## Список литературы

1. Иванов И. И. Книга. — М., 2020.
2. Петров П. П. Статья. — СПб., 2021.

[^1]: Текст сноски.
"""


def convert(tmp_path, text=SAMPLE, **options):
    src = tmp_path / "in.md"
    src.write_text(text, encoding="utf-8")
    settings = DocumentSettings()
    settings.auto_numbering_headings = True
    for key, value in options.items():
        setattr(settings, key, value)
    out = tmp_path / "out.docx"
    MarkdownToDocxConverter(settings).convert(str(src), str(out))
    return out


def texts(path):
    return [p.text for p in Document(str(path)).paragraphs if p.text.strip()]


def test_page_number_field_and_title_page(tmp_path):
    out = convert(tmp_path)
    with zipfile.ZipFile(out) as z:
        footers = [z.read(n).decode() for n in z.namelist() if n.startswith("word/footer")]
        document = z.read("word/document.xml").decode()
    assert any("PAGE" in f for f in footers)
    assert "<w:titlePg/>" in document


def test_heading_font_is_not_theme_font(tmp_path):
    out = convert(tmp_path)
    with zipfile.ZipFile(out) as z:
        styles = z.read("word/styles.xml").decode()
    heading = styles[styles.index('w:styleId="Heading1"'):]
    heading = heading[:heading.index("</w:style>")]
    assert "asciiTheme" not in heading and 'w:ascii="Times New Roman"' in heading


def test_lists_have_single_dash_nesting_and_restart(tmp_path):
    lines = texts(convert(tmp_path))
    assert "– Первый пункт" in lines
    nested = next(p for p in Document(str(convert(tmp_path))).paragraphs if p.text == "– Вложенный пункт")
    assert nested.paragraph_format.left_indent.cm > 0
    assert lines.count("1) Раз") == 1 and "1) Снова один" in lines


def test_structural_headings_are_not_numbered(tmp_path):
    lines = texts(convert(tmp_path))
    assert "ВВЕДЕНИЕ" in lines and "СПИСОК ЛИТЕРАТУРЫ" in lines
    assert "1.1. Цели" in lines  # нумерация разделов идёт мимо «Введения»
    assert "1. Иванов И. И. Книга. — М., 2020." in lines


def test_table_and_pipe_paragraph(tmp_path):
    out = convert(tmp_path)
    doc = Document(str(out))
    assert len(doc.tables) == 1 and doc.tables[0].cell(1, 1).text == "1"
    assert "Строка с | вертикальной чертой, но не таблица." in texts(out)


def test_cli(tmp_path, capsys):
    src = tmp_path / "doc.md"
    src.write_text("# Заголовок\n\nТекст\n", encoding="utf-8")
    assert cli_main([str(src), "--no-heading-numbers"]) == 0
    out = tmp_path / "doc.docx"
    assert out.exists() and "Заголовок" in texts(out)
    assert cli_main([str(tmp_path / "nope.md")]) == 1
